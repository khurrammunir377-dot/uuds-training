import os
import re
import sqlite3
import openpyxl
from datetime import datetime, date, timedelta
from database import get_db, init_db

COURSES_CATALOG = [
    {"code": "OL-1237", "name": "SMS for All", "category": "Safety Management", "validity_months": 24, "desc": "Emirates Group SMS training for all personnel"},
    {"code": "OL-723", "name": "Fundamentals of Human Factors", "category": "Human Factors", "validity_months": 24, "desc": "Human Factors fundamentals in aviation maintenance"},
    {"code": "OL-1666", "name": "EASA Regulations - Initial/Continuation training", "category": "Regulations", "validity_months": 24, "desc": "EASA Part 145 regulatory requirements"},
    {"code": "OL-1667", "name": "UAE GCAA Aviation Legislation", "category": "Regulations", "validity_months": 24, "desc": "UAE GCAA CAR 145 aviation legislation"},
    {"code": "OL-1653", "name": "EWIS 3, 4 & 5 - Initial/Continuation training", "category": "Technical", "validity_months": 24, "desc": "Electrical Wiring Interconnection System"},
    {"code": "OL-1073", "name": "Reduced Vertical Separation Minimum (RVSM)", "category": "Technical", "validity_months": 24, "desc": "RVSM operational and maintenance compliance"},
    {"code": "OL-1678", "name": "Fuel Tank Safety Phase 1 and Phase 2", "category": "Technical", "validity_months": 24, "desc": "Fuel Tank Safety Phase 1 and 2 regulations"},
    {"code": "OL-1707", "name": "Emirates Engineering Procedures", "category": "Procedures", "validity_months": 24, "desc": "EK SOPs and department working procedures"},
    {"code": "OL-2550", "name": "Working at Heights", "category": "Safety", "validity_months": 24, "desc": "Aviation ramp and hangar working at heights safety"},
    {"code": "OL-1395", "name": "An Introduction to Cyber Security", "category": "IT & Security", "validity_months": 12, "desc": "Aviation cybersecurity awareness and protocols"},
    {"code": "OL-1732", "name": "Data Privacy Basics", "category": "IT & Security", "validity_months": 12, "desc": "General data protection and Emirates privacy policy"},
    {"code": "OL-2136", "name": "Communicable Disease", "category": "Health & Safety", "validity_months": 24, "desc": "Health precautions and contagious illness awareness"},
    {"code": "OL-2529", "name": "The Emirates Group Safety Policy and You", "category": "Safety Management", "validity_months": 24, "desc": "Corporate safety culture, obligations and reporting"},
    {"code": "OL-1968", "name": "ETOPS Generic and Type Initial/Continuation", "category": "Technical", "validity_months": 24, "desc": "Extended-range Twin-engine Operational Performance"},
    {"code": "OL-1965", "name": "Workshop Procedures", "category": "Procedures", "validity_months": 24, "desc": "Workshop tooling, safety and maintenance documentation"},
    {"code": "OL-1090", "name": "Emirates Airside Safety", "category": "Safety", "validity_months": 24, "desc": "Airside driving, ramp etiquette, and safety guidelines"},
    {"code": "OL-498", "name": "Continuation Training in Safety and Human Factors", "category": "Human Factors", "validity_months": 24, "desc": "Recurrent 24-month refresher following initial HF"},
    {"code": "OL-2669", "name": "B777 Retrofit - Falcon Project", "category": "Aircraft Specific", "validity_months": 24, "desc": "Falcon Project Boeing 777 cabin reconfiguration"},
    {"code": "OL-2208", "name": "Data Privacy Basics - Refresher", "category": "IT & Security", "validity_months": 12, "desc": "Annual refresher for Data Privacy Basics"},
    {"code": "GEN-101", "name": "Induction Safety Briefing", "category": "Onboarding", "validity_months": 36, "desc": "Company induction, hangar safety, security orientation"},
    {"code": "GEN-102", "name": "Competency Assessment", "category": "Assessment", "validity_months": 24, "desc": "Technical skills and authorization competency assessment"},
]

def parse_date_value(val):
    if val is None:
        return None, None
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d"), None
    
    s = str(val).strip()
    if not s:
        return None, None
    
    # Fix known typos in dataset
    s_fixed = re.sub(r'(\d{1,2})-([A-Za-z]+)-207$', r'\1-\2-2027', s)
    s_fixed = re.sub(r'(\d{1,2})([A-Za-z]+)-(\d{4})', r'\1-\2-\3', s_fixed)
    
    date_formats = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d-%b-%Y",
        "%d-%b-%y",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d-%m-%Y"
    ]
    for fmt in date_formats:
        try:
            d = datetime.strptime(s_fixed, fmt)
            # Normalize 2-digit years
            if d.year < 100:
                d = d.replace(year=d.year + 2000)
            return d.strftime("%Y-%m-%d"), None
        except Exception:
            pass
            
    return None, s

def calculate_status(expiry_date_str, text_note=None, ref_date=None):
    if text_note:
        t = text_note.lower().strip()
        if "yes" in t:
            return "Valid"
        if "planned" in t:
            return "Not Recorded"
            
    if not expiry_date_str:
        return "Not Recorded"
        
    if ref_date is None:
        ref_date = date.today()
        
    try:
        exp = datetime.strptime(expiry_date_str, "%Y-%m-%d").date()
        diff = (exp - ref_date).days
        if diff < 0:
            return "Overdue"
        elif diff <= 30:
            return "Due Within 30 Days"
        else:
            return "Valid"
    except Exception:
        return "Not Recorded"

def match_column_to_course_code(header_text):
    if not header_text:
        return None
    h = str(header_text).lower()
    
    if "easa" in h and "145" in h:
        return "OL-1666"
    if "gcaa" in h and "145" in h:
        return "OL-1667"
    if "rvsm" in h:
        return "OL-1073"
    if "sop" in h or "ek procedures" in h:
        return "OL-1707"
    if "sms" in h:
        return "OL-1237"
    if "fuel tank" in h:
        return "OL-1678"
    if "ewis" in h:
        return "OL-1653"
    if "cyber" in h:
        return "OL-1395"
    if "data privacy" in h:
        return "OL-1732"
    if "height" in h:
        return "OL-2550"
    if "fundamental" in h and "human" in h:
        return "OL-723"
    if ("contiuat" in h or "continuation" in h) and ("safety" in h or "hf" in h):
        return "OL-498"
    if "induction" in h:
        return "GEN-101"
    if "competency" in h:
        return "GEN-102"
    if "etops" in h:
        return "OL-1968"
    if "workshop" in h:
        return "OL-1965"
    if "airside" in h:
        return "OL-1090"
    if "retrofit" in h or "falcon" in h:
        return "OL-2669"
        
    return None

def import_all_data(excel_path=None, mobile_excel_path=None):
    init_db()
    conn = get_db()
    cursor = conn.cursor()
    
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    if not excel_path:
        excel_path = os.path.join(base_dir, "Manpower and Training Schedule- R4 27 Aug 2026 (version 1) - Copy.xlsx")
    if not mobile_excel_path:
        mobile_excel_path = os.path.join(base_dir, "staff with mobile no.xlsx")
        
    print(f"Loading courses catalog ({len(COURSES_CATALOG)} courses)...")
    for c in COURSES_CATALOG:
        cursor.execute("""
        INSERT INTO courses (code, name, category, validity_months, description)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(code) DO UPDATE SET
            name = excluded.name,
            category = excluded.category,
            validity_months = excluded.validity_months,
            description = excluded.description
        """, (c["code"], c["name"], c["category"], c["validity_months"], c["desc"]))
    conn.commit()
    
    # Load course IDs
    cursor.execute("SELECT id, code FROM courses")
    course_map = {row["code"]: row["id"] for row in cursor.fetchall()}
    
    # Load mobile numbers and complete staff info lookup
    mobile_dict = {}
    staff_lookup_dict = {}
    if os.path.exists(mobile_excel_path):
        print(f"Reading mobile numbers and staff details from {mobile_excel_path}...")
        wb_m = openpyxl.load_workbook(mobile_excel_path, data_only=True)
        ws_m = wb_m["STAFF"] if "STAFF" in wb_m.sheetnames else wb_m.active
        for r in range(2, ws_m.max_row + 1):
            raw_uuds = str(ws_m.cell(r, 1).value or "").strip().upper()
            clean_uuds = raw_uuds.replace("_", "-")
            con_no = str(ws_m.cell(r, 2).value or "").strip()
            name_val = str(ws_m.cell(r, 3).value or "").strip()
            team_val = str(ws_m.cell(r, 4).value or "").strip()
            phone_val = str(ws_m.cell(r, 5).value or "").strip()
            if clean_uuds:
                mobile_dict[clean_uuds] = phone_val
                mobile_dict[raw_uuds] = phone_val
                info = {
                    "contingent_id": con_no,
                    "full_name": name_val,
                    "team": team_val,
                    "mobile_no": phone_val
                }
                staff_lookup_dict[clean_uuds] = info
                staff_lookup_dict[raw_uuds] = info
        print(f"Loaded {len(mobile_dict)} mobile numbers & staff profiles from contact list.")
    else:
        print("Mobile Excel file not found, continuing without it.")

    # Read training sheets
    if not os.path.exists(excel_path):
        print(f"Error: Training schedule file not found at {excel_path}")
        return
        
    print(f"Reading training workbook from {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    sheet_status_map = {
        "Training Register": "Active",
        "UN Paid Leave": "UN Paid Leave",
        "DWC TX": "DWC TX",
        "Non Tec.": "Non Tec.",
        "RESIGN": "RESIGN",
        "Redundancy": "Redundancy"
    }
    
    total_employees_imported = 0
    total_records_imported = 0
    
    for sheet_name in wb.sheetnames:
        emp_status = sheet_status_map.get(sheet_name)
        if not emp_status:
            continue
            
        ws = wb[sheet_name]
        print(f"\nProcessing Sheet: '{sheet_name}' (Status: {emp_status})...")
        
        header_cols = {}
        col_uuds = None
        col_staff = None
        col_name = None
        col_pos = None
        col_team = None
        col_start = None
        start_row = 2

        if sheet_name == "Redundancy":
            col_pos = 1
            col_team = 3
            col_name = 4
            col_start = 5
            col_staff = 6
            col_uuds = 7
            start_row = 2
            # 12 core courses from column 8
            header_cols = {
                8: "OL-1666",
                9: "OL-1667",
                10: "OL-1073",
                11: "OL-1707",
                12: "OL-1237",
                13: "OL-1678",
                14: "OL-1653",
                15: "OL-1395",
                16: "OL-1732",
                17: "OL-2550",
                18: "OL-723",
                19: "OL-498"
            }
        elif sheet_name == "Non Tec.":
            col_name = 2
            col_staff = 3
            col_uuds = 4
            col_start = 5
            col_team = 6
            start_row = 5
            header_cols = {
                7: "OL-1666",
                8: "OL-1667",
                9: "OL-1073"
            }
        else:
            # Scan header row (Row 1)
            for c in range(1, ws.max_column + 1):
                val = ws.cell(1, c).value
                if val:
                    val_str = str(val).strip()
                    course_code = match_column_to_course_code(val_str)
                    if course_code:
                        header_cols[c] = course_code
                        
            # Look in row 1 for metadata headers
            for c in range(1, ws.max_column + 1):
                h = str(ws.cell(1, c).value or "").strip().lower()
                if "uuds" in h:
                    col_uuds = c
                elif "staff" in h or "contingent" in h:
                    col_staff = c
                elif "name" in h and "company" not in h:
                    col_name = c
                elif "position" in h:
                    col_pos = c
                elif "team" in h:
                    col_team = c
                elif "start date" in h or "dxb start" in h:
                    col_start = c
                    
        print(f"Sheet '{sheet_name}': Found {len(header_cols)} training courses mapped to headers.")
        
        # Iterate data rows
        for r in range(start_row, ws.max_row + 1):
            # Try to get UUDS No
            uuds_val = ws.cell(r, col_uuds).value if col_uuds else None
            name_val = ws.cell(r, col_name).value if col_name else None
            
            # If UUDS is missing but name exists, scan row
            if not uuds_val:
                for c in range(1, 15):
                    v = str(ws.cell(r, c).value or "").strip()
                    if re.match(r"^UUDS-\d+", v, re.IGNORECASE):
                        uuds_val = v
                        break
                        
            if not uuds_val and not name_val:
                continue # empty row
                
            uuds_str = str(uuds_val).strip().upper().replace("_", "-") if uuds_val else ""
            if not uuds_str:
                # If no UUDS number, generate a synthetic unique identifier
                uuds_str = f"EXT-{sheet_name[:3].upper()}-{r:03d}"
                
            name_str = str(name_val).strip() if name_val else ""
            staff_no_str = str(ws.cell(r, col_staff).value or "").strip() if col_staff else ""
            pos_str = str(ws.cell(r, col_pos).value or "").strip() if col_pos else ""
            team_str = str(ws.cell(r, col_team).value or "").strip() if col_team else ""

            # Check fallback in staff_lookup_dict
            matched_staff = staff_lookup_dict.get(uuds_str, {})
            if not name_str or name_str.lower() in ["staff member", "none"]:
                if matched_staff.get("full_name"):
                    name_str = matched_staff["full_name"]
                else:
                    name_str = "Staff Member"

            if not staff_no_str and matched_staff.get("contingent_id"):
                staff_no_str = matched_staff["contingent_id"]

            if not team_str and matched_staff.get("team"):
                team_str = matched_staff["team"]
            
            # Special check for Malith Rangana in UN Paid Leave
            if "malith rangana" in name_str.lower():
                uuds_str = "UUDS-1488"
                staff_no_str = "C74497"
                if not pos_str:
                    pos_str = "Assistant Store Keeper"
                if not team_str:
                    team_str = "Stores"
            
            # Start date
            start_date_str = None
            if col_start:
                s_d, _ = parse_date_value(ws.cell(r, col_start).value)
                start_date_str = s_d
                
            # Mobile number from mobile dictionary or lookup
            mobile_str = mobile_dict.get(uuds_str, "")
            if not mobile_str and matched_staff.get("mobile_no"):
                mobile_str = matched_staff["mobile_no"]
            if not mobile_str and "malith rangana" in name_str.lower():
                mobile_str = "0522375921"
            
            # Insert or update employee
            cursor.execute("""
            INSERT INTO employees (uuds_no, contingent_id, full_name, position, team, mobile_no, dxb_start_date, employment_status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(uuds_no) DO UPDATE SET
                contingent_id = coalesce(nullif(excluded.contingent_id, ''), employees.contingent_id),
                full_name = excluded.full_name,
                position = coalesce(nullif(excluded.position, ''), employees.position),
                team = coalesce(nullif(excluded.team, ''), employees.team),
                mobile_no = coalesce(nullif(excluded.mobile_no, ''), employees.mobile_no),
                dxb_start_date = coalesce(excluded.dxb_start_date, employees.dxb_start_date),
                employment_status = excluded.employment_status,
                updated_at = CURRENT_TIMESTAMP
            """, (uuds_str, staff_no_str, name_str, pos_str, team_str, mobile_str, start_date_str, emp_status))
            
            # Get employee DB ID
            cursor.execute("SELECT id FROM employees WHERE uuds_no = ?", (uuds_str,))
            emp_id = cursor.fetchone()["id"]
            total_employees_imported += 1
            
            # Insert training records for this employee
            for c, code in header_cols.items():
                cell_val = ws.cell(r, c).value
                date_parsed, text_note = parse_date_value(cell_val)
                c_id = course_map.get(code)
                if not c_id:
                    continue
                    
                status = calculate_status(date_parsed, text_note)
                
                completion_dt = None
                expiry_dt = date_parsed
                
                cursor.execute("""
                INSERT INTO training_records (employee_id, course_id, completion_date, expiry_date, status, notes, updated_by)
                VALUES (?, ?, ?, ?, ?, ?, 'excel_importer')
                ON CONFLICT(employee_id, course_id) DO UPDATE SET
                    expiry_date = excluded.expiry_date,
                    status = excluded.status,
                    notes = coalesce(excluded.notes, training_records.notes),
                    updated_at = CURRENT_TIMESTAMP
                """, (emp_id, c_id, completion_dt, expiry_dt, status, text_note))
                total_records_imported += 1
                
        conn.commit()

    # Ensure all employees (active, unpaid leave, dwc tx, etc.) have entries for all courses
    print("\nEnsuring all courses are represented for each employee...")
    cursor.execute("SELECT id FROM employees")
    all_emp_ids = [row["id"] for row in cursor.fetchall()]
    
    cursor.execute("SELECT id FROM courses")
    all_course_ids = [row["id"] for row in cursor.fetchall()]
    
    for emp_id in all_emp_ids:
        for c_id in all_course_ids:
            cursor.execute("""
            INSERT OR IGNORE INTO training_records (employee_id, course_id, status, notes, updated_by)
            VALUES (?, ?, 'Not Recorded', 'Initial requirement', 'system')
            """, (emp_id, c_id))
            
    conn.commit()
    conn.close()
    
    print("\nImport completed successfully!")
    print(f"- Total employees imported/synced: {total_employees_imported}")
    print(f"- Total training compliance records: {total_records_imported}")

if __name__ == "__main__":
    import_all_data()
